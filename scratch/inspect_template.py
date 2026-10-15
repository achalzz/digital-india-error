import docx

doc = docx.Document('Milestone_I_Report_Template.docx')

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

# Iterate over body elements in order
body_elements = []
for child in doc.element.body:
    if child.tag.endswith('p'):
        p = docx.text.paragraph.Paragraph(child, doc)
        text = p.text.strip()
        if text:
            body_elements.append(('P', p.style.name, text[:80]))
    elif child.tag.endswith('tbl'):
        t = docx.table.Table(child, doc)
        header = [c.text.strip().replace('\n', ' ') for c in t.rows[0].cells] if len(t.rows) > 0 else []
        body_elements.append(('TBL', f"{len(t.rows)}x{len(t.columns)}", " | ".join(header)[:80]))

for idx, (kind, meta, preview) in enumerate(body_elements):
    print(f"{idx:02d} [{kind}] ({meta}): {preview}")
